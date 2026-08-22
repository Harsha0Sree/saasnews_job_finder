#!/usr/bin/env python3
"""
Combine multiple saasnews_scraper.py output xlsx files into one master file.
De-duplicates by job URL, applies the latest location filter, and produces a
clean master workbook with the best available data per match.
"""
import glob
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---- Filter patterns (mirrors saasnews_scraper.py) ----
RESTRICTED_REMOTE_REJECT = [
    r"remote\s*\(us\b", r"remote\s*\(united\s+states", r"remote\s*\(usa?\)",
    r"remote\s*\(uk\b", r"remote\s*\(united\s+kingdom",
    r"remote\s*\(eu\)", r"remote\s*\(europe\)", r"remote\s*\(emea\)",
    r"remote\s*\(germany\)", r"remote\s*\(france\)",
    r"remote\s*\(canada\)", r"remote\s*\(aus\w*\)",
    r"remote\s*\(latam\)", r"remote\s*\(apac\)",
    r"us\s+only", r"uk\s+only", r"eu\s+only",
    r"north\s+america\s+only", r"europe\s+only",
    r"^\s*u\.?s\.?a?\.?\s*(?:,|/|\|)?\s*(?:remote|hybrid)?\s*$",
    r"^\s*u\.?k\.?\s*(?:,|/|\|)?\s*(?:remote|hybrid)?\s*$",
    r"^\s*eu\.?\s*(?:,|/|\|)?\s*(?:remote|hybrid)?\s*$",
    r"\bus\s*,\s*remote\b",
    r"\buk\s*,\s*remote\b",
    r"\beu\s*,\s*remote\b",
    r"\bunited\s+states\s*,\s*remote\b",
    r"\bunited\s+kingdom\s*,\s*remote\b",
    r"\bremote\s*,\s*u\.?s\.?a?\.?\b",
    r"\bremote\s*,\s*united\s+states\b",
    r"\bremote\s*,\s*u\.?k\.?\b",
    r"\bremote\s*,\s*united\s+kingdom\b",
    r"\bremote\s*,\s*eu\.?\b",
    r"\bremote\s*,\s*europe\b",
    r"\bremote\s*,\s*canada\b",
    r"\bsan\s+francisco\s*/\s*remote\b",
    r"\bnew\s+york\s*/\s*remote\b",
]
RESTRICTED_RE = re.compile("|".join(RESTRICTED_REMOTE_REJECT), re.IGNORECASE)

# US city names — if a location mentions these AND "remote" without an India
# token, it's effectively US-restricted remote (not worldwide).
US_CITY_RE = re.compile(
    r"\b(san\s+francisco|new\s+york|seattle|austin|boston|chicago|"
    r"los\s+angeles|denver|atlanta|portland|washington\s+dc|"
    r"palo\s+alto|mountain\s+view|menlo\s+park|redwood\s+city|"
    r"cambridge|irvine|dallas|houston|phoenix|minneapolis|"
    r"sf\b|nyc\b|la\b|dc\b)\b",
    re.IGNORECASE,
)
# UK/EU cities — same logic.
EU_CITY_RE = re.compile(
    r"\b(london|berlin|paris|amsterdam|dublin|madrid|barcelona|"
    r"munich|frankfurt|vienna|stockholm|copenhagen|zurich|"
    r"brussels|lisbon|prague|warsaw|helsinki|oslo)\b",
    re.IGNORECASE,
)
# Reject locations that look like Python dict reprs (from broken JSON-LD
# parsing in earlier runs).
DICT_REPR_RE = re.compile(r"^\s*\{.*\}\s*$", re.DOTALL)

INDIA_TOKENS = [
    "india", "bangalore", "bengaluru", "mumbai", "delhi", "new delhi",
    "noida", "gurugram", "gurgaon", "pune", "hyderabad", "chennai",
    "kolkata", "ahmedabad", "jaipur", "kochi", "coimbatore", "indore",
    "chandigarh", "lucknow", "bhubaneswar", "trivandrum", "thiruvananthapuram",
    "visakhapatnam", "remote.*india", "india.*remote",
]
INDIA_RE = re.compile("|".join(INDIA_TOKENS), re.IGNORECASE)
BARE_REMOTE = re.compile(r"\bremote\b", re.IGNORECASE)

# ---- Styling ----
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
HIGH_FILL = PatternFill("solid", fgColor="D1FAE5")
MED_FILL = PatternFill("solid", fgColor="FEF3C7")
LOW_FILL = PatternFill("solid", fgColor="FEE2E2")


def location_passes(loc: str) -> bool:
    """Return True if the location passes the India-or-Remote-Worldwide filter."""
    if not loc:
        return False
    # Reject dict reprs from broken JSON-LD parsing.
    if DICT_REPR_RE.search(loc):
        return False
    if RESTRICTED_RE.search(loc):
        return False
    if INDIA_RE.search(loc):
        return True
    # If location mentions US/EU cities AND "remote", treat as restricted.
    if BARE_REMOTE.search(loc) and (US_CITY_RE.search(loc) or EU_CITY_RE.search(loc)):
        return False
    if BARE_REMOTE.search(loc):
        return True
    return False


def clean_location(loc: str) -> str:
    """Tidy up a location string: extract the most relevant city/remote token."""
    if not loc:
        return ""
    # 1. If an India city is mentioned, extract just that.
    india_cities = [
        "bengaluru", "bangalore", "mumbai", "delhi", "new delhi", "noida",
        "gurugram", "gurgaon", "pune", "hyderabad", "chennai", "kolkata",
        "ahmedabad", "jaipur", "kochi", "coimbatore", "indore",
        "chandigarh", "lucknow", "bhubaneswar", "trivandrum",
        "thiruvananthapuram", "visakhapatnam",
    ]
    for city in india_cities:
        m = re.search(r"\b" + city + r"\b", loc, re.IGNORECASE)
        if m:
            c = m.group(0)
            # Normalize Bangalore → Bengaluru
            if c.lower() == "bangalore":
                c = "Bengaluru"
            elif c.lower() == "new delhi":
                c = "New Delhi"
            else:
                c = c.capitalize()
            return c + ", India"
    # 2. If "Remote" is in the text, extract the remote qualifier.
    m = re.search(r"\bremote\b", loc, re.IGNORECASE)
    if m:
        # Look for a parenthetical qualifier after "remote".
        after = loc[m.end():m.end() + 30]
        qm = re.match(r"\s*\(([^)]+)\)", after)
        if qm:
            qualifier = qm.group(1).strip()
            return f"Remote ({qualifier})"
        return "Remote"
    # 3. Strip common garbage and return truncated.
    parts = re.split(r"[·•|]+", loc)
    if len(parts) > 1:
        for p in parts:
            p = p.strip()
            if p and len(p) < 60 and re.search(r"[A-Z][a-z]+", p):
                return p
    return loc.strip(" ·,|•\u2022-—")[:60]


def main():
    out_dir = Path("/home/mikeysama/products/saasnews-job-finder/download")
    files = sorted(glob.glob(str(out_dir / "saasnews_jobs_*.xlsx")))
    if not files:
        print("No input files found.")
        sys.exit(1)
    print(f"Combining {len(files)} files:")
    for f in files:
        print(f"  - {f}")

    # Matches: dedup by job_url, keep the cleanest location.
    # Companies: dedup by news_url.
    matches_by_url: dict[str, dict] = {}
    companies_by_news: dict[str, dict] = {}
    errors: list[dict] = []
    stats = {"files": 0, "matches_in": 0, "matches_dedup": 0,
             "matches_filtered": 0, "companies_in": 0, "companies_dedup": 0}

    for f in files:
        stats["files"] += 1
        wb = load_workbook(f, read_only=True)

        # Matches
        if "Matches" in wb.sheetnames:
            ws = wb["Matches"]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                stats["matches_in"] += 1
                rec = dict(zip(headers, row))
                url = rec.get("Job URL") or ""
                if not url:
                    continue
                loc = rec.get("Location") or ""
                # Re-apply the (updated) location filter.
                if not location_passes(loc):
                    continue
                # Dedup: keep the record with the cleanest location.
                existing = matches_by_url.get(url)
                if existing is None:
                    matches_by_url[url] = rec
                    stats["matches_dedup"] += 1
                else:
                    # Prefer the one with the shorter (cleaner) location.
                    if len(loc) < len(existing.get("Location") or ""):
                        matches_by_url[url] = rec

        # All Companies
        if "All Companies" in wb.sheetnames:
            ws = wb["All Companies"]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                stats["companies_in"] += 1
                rec = dict(zip(headers, row))
                news_url = rec.get("News URL") or ""
                if not news_url:
                    continue
                existing = companies_by_news.get(news_url)
                if existing is None:
                    companies_by_news[news_url] = rec
                    stats["companies_dedup"] += 1
                else:
                    # Prefer the one with a careers page found / more jobs.
                    if (rec.get("Careers Status") == "found" and
                            existing.get("Careers Status") != "found"):
                        companies_by_news[news_url] = rec
                    elif (rec.get("Jobs Matched", 0) or 0) > (existing.get("Jobs Matched", 0) or 0):
                        companies_by_news[news_url] = rec

        # Errors
        if "Errors" in wb.sheetnames:
            ws = wb["Errors"]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rec = dict(zip(headers, row))
                errors.append(rec)
        wb.close()

    stats["matches_filtered"] = len(matches_by_url)
    print(f"\nStats: {stats}")

    # Clean up match locations.
    for url, rec in matches_by_url.items():
        rec["Location"] = clean_location(rec.get("Location") or "")

    # ---- Write master workbook ----
    wb = Workbook()
    wb.properties.creator = "Z.ai"

    # --- Sheet 1: Matches ---
    ws1 = wb.active
    ws1.title = "Matches"
    headers1 = [
        "Company", "Job Title", "Category", "Confidence", "Matched Keyword",
        "Location", "Location Basis", "Fresher Basis", "Fresher Level",
        "Apply URL", "Job URL",
        "Company Website", "Careers Page", "Funding Round", "Funding Date",
        "Software Category", "News Headline", "News URL",
    ]
    for c, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

    rows = []
    for rec in matches_by_url.values():
        rows.append([
            rec.get("Company", ""), rec.get("Job Title", ""),
            rec.get("Category", ""), rec.get("Confidence", ""),
            rec.get("Matched Keyword", ""), rec.get("Location", ""),
            rec.get("Location Basis", ""),
            rec.get("Fresher Basis", ""), rec.get("Fresher Level", ""),
            rec.get("Apply URL", ""), rec.get("Job URL", ""),
            rec.get("Company Website", ""), rec.get("Careers Page", ""),
            rec.get("Funding Round", ""), rec.get("Funding Date", ""),
            rec.get("Software Category", ""), rec.get("News Headline", ""),
            rec.get("News URL", ""),
        ])
    conf_order = {"high": 0, "medium": 1, "low": 2}
    rows.sort(key=lambda r: (conf_order.get(str(r[3]).lower(), 9),
                              0 if r[2] == "AI/ML" else 1 if r[2] == "Backend" else 2,
                              r[0].lower()))
    color_map = {"high": HIGH_FILL, "medium": MED_FILL, "low": LOW_FILL}
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if c_idx == 4:  # Confidence column
                fill = color_map.get(str(val).lower())
                if fill:
                    cell.fill = fill
    widths = [22, 42, 12, 12, 20, 28, 18, 22, 14, 50, 50, 28, 38, 14, 14, 22, 50, 50]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A2"
    if rows:
        ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers1))}{len(rows) + 1}"

    # --- Sheet 2: All Companies ---
    ws2 = wb.create_sheet("All Companies")
    headers2 = [
        "Company", "Website", "Funding Round", "Funding Date", "Lead Investor",
        "Software Category", "Careers Page", "Careers Status",
        "Jobs Found", "Jobs Matched", "News Headline", "News URL", "Error",
    ]
    for c, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    comp_rows = []
    for rec in companies_by_news.values():
        comp_rows.append([
            rec.get("Company", ""), rec.get("Website", ""),
            rec.get("Funding Round", ""), rec.get("Funding Date", ""),
            rec.get("Lead Investor", ""), rec.get("Software Category", ""),
            rec.get("Careers Page", ""), rec.get("Careers Status", ""),
            rec.get("Jobs Found", 0) or 0, rec.get("Jobs Matched", 0) or 0,
            rec.get("News Headline", ""), rec.get("News URL", ""),
            rec.get("Error", ""),
        ])
    comp_rows.sort(key=lambda r: (r[7] != "found", -int(r[9] or 0), r[0].lower()))
    for r_idx, row in enumerate(comp_rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    widths2 = [22, 30, 16, 14, 22, 24, 38, 14, 12, 12, 50, 50, 40]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"
    if comp_rows:
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{len(comp_rows) + 1}"

    # --- Sheet 3: Errors ---
    ws3 = wb.create_sheet("Errors")
    headers3 = ["Company", "Website", "News URL", "Stage", "Error", "Timestamp"]
    for c, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    err_rows = [[e.get("Company", ""), e.get("Website", ""), e.get("News URL", ""),
                 e.get("Stage", ""), e.get("Error", ""), e.get("Timestamp", "")]
                for e in errors]
    for r_idx, row in enumerate(err_rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    widths3 = [22, 30, 50, 16, 60, 26]
    for i, w in enumerate(widths3, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"

    # --- Sheet 4: Run Summary ---
    ws4 = wb.create_sheet("Run Summary")
    conf_counts = {"high": 0, "medium": 0, "low": 0}
    for r in rows:
        conf_counts[r[3]] = conf_counts.get(r[3], 0) + 1
    cat_counts = {}
    for r in rows:
        cat_counts[r[2]] = cat_counts.get(r[2], 0) + 1
    summary_rows = [
        ["Master File Generated", datetime.now(timezone.utc).isoformat()],
        ["Source Files Combined", stats["files"]],
        ["Total News Articles (dedup)", len(companies_by_news)],
        ["Companies With Careers Page", sum(1 for c in companies_by_news.values() if c.get("Careers Status") == "found")],
        ["Companies With No Careers Page", sum(1 for c in companies_by_news.values() if c.get("Careers Status") == "not_found")],
        ["Total Errors Logged", len(errors)],
        ["Total Matching Jobs (after dedup + filter)", len(rows)],
        ["  High Confidence", conf_counts.get("high", 0)],
        ["  Medium Confidence", conf_counts.get("medium", 0)],
        ["  Low Confidence", conf_counts.get("low", 0)],
        ["Matches by Category", ""],
        ["  AI/ML", cat_counts.get("AI/ML", 0)],
        ["  Backend", cat_counts.get("Backend", 0)],
        ["  Full-stack", cat_counts.get("Full-stack", 0)],
    ]
    for r, (k, v) in enumerate(summary_rows, 1):
        ws4.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws4.cell(row=r, column=2, value=v)
    ws4.column_dimensions["A"].width = 42
    ws4.column_dimensions["B"].width = 80

    out_path = out_dir / "saasnews_jobs_MASTER.xlsx"
    wb.save(out_path)
    print(f"\nMaster file: {out_path}")
    print(f"  Unique companies: {len(companies_by_news)}")
    print(f"  Unique matching jobs: {len(rows)}")
    print(f"    high: {conf_counts.get('high', 0)}  medium: {conf_counts.get('medium', 0)}  low: {conf_counts.get('low', 0)}")


if __name__ == "__main__":
    main()
