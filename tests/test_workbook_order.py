"""Workbook ordering + tracking columns.

Seam under test: _write_xlsx_from_checkpoint — the Matches sheet must read
newest-scraped-first ("First Seen" descending), and carry user tracking
columns (Applied / Feedback / Notes) with dropdown validation.
"""
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from saasnews_scraper import JobMatch, _append_checkpoint, _load_checkpoint, \
    _write_xlsx_from_checkpoint

NEW_HEADERS = ["First Seen", "Applied", "Feedback", "Notes"]


def _match(i: int, scraped_at: str) -> JobMatch:
    return JobMatch(
        company_name=f"Company{i}",
        company_website="https://co.example.com",
        news_url="https://example.com/news/1/",
        news_headline="Company raises",
        funding_round="Seed",
        funding_date="July 2026",
        software_category="AI",
        job_title="Python Backend Engineer",
        job_url=f"https://co.example.com/jobs/{i}",
        location="Bengaluru",
        apply_url=f"https://co.example.com/jobs/{i}",
        matched_keyword="python",
        match_category="Backend",
        confidence="high",
        location_basis="india",
        careers_page_url="https://co.example.com/careers",
        posted_date="",
        fresher_basis="title_positive",
        fresher_level="explicit",
        scraped_at=scraped_at,
    )


def _seed(tmp_path, matches):
    out = str(tmp_path)
    record = {
        "news_url": "https://example.com/news/1/",
        "company_rec": {
            "company_name": "Company", "company_website": "https://co.example.com",
            "news_url": "https://example.com/news/1/", "news_headline": "",
            "funding_round": "", "funding_date": "", "lead_investor": "",
            "software_category": "", "careers_page_url": "",
            "careers_page_status": "found", "jobs_found": len(matches),
            "jobs_matched": len(matches), "error": "",
        },
        "matches": [
            {
                "company_name": m.company_name,
                "company_website": m.company_website,
                "news_url": m.news_url,
                "news_headline": m.news_headline,
                "funding_round": m.funding_round,
                "funding_date": m.funding_date,
                "software_category": m.software_category,
                "job_title": m.job_title,
                "job_url": m.job_url,
                "location": m.location,
                "apply_url": m.apply_url,
                "matched_keyword": m.matched_keyword,
                "match_category": m.match_category,
                "confidence": m.confidence,
                "location_basis": m.location_basis,
                "careers_page_url": m.careers_page_url,
                "posted_date": m.posted_date,
                "fresher_basis": m.fresher_basis,
                "fresher_level": m.fresher_level,
                "scraped_at": m.scraped_at,
            }
            for m in matches
        ],
        "error": None,
    }
    _append_checkpoint(out, record)
    return out, _load_checkpoint(out)


def base_time():
    return datetime(2026, 8, 24, 9, 0, 0, tzinfo=timezone.utc)


def test_matches_sorted_newest_first_with_first_seen_column(tmp_path):
    t = base_time()
    older = _match(1, (t - timedelta(days=2)).isoformat())
    newer = _match(2, t.isoformat())
    out, ckpt = _seed(tmp_path, [older, newer])

    path = _write_xlsx_from_checkpoint(out, ckpt, {})
    wb = load_workbook(path)
    ws = wb["Matches"]
    headers = [c.value for c in ws[1]]
    for h in NEW_HEADERS:
        assert h in headers
    first_seen_col = headers.index("First Seen") + 1
    job_col = headers.index("Job URL") + 1

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    # Newest scraped job on top.
    assert rows[0][job_col - 1] == "https://co.example.com/jobs/2"
    assert rows[1][job_col - 1] == "https://co.example.com/jobs/1"
    # First Seen values present and in descending order.
    assert rows[0][first_seen_col - 1] > rows[1][first_seen_col - 1]
    wb.close()


def test_matches_without_timestamp_sink_to_bottom(tmp_path):
    """Old checkpoints have no scraped_at — they must sort below timestamped
    rows, never crash."""
    t = base_time()
    stampless = _match(3, "")
    fresh = _match(4, t.isoformat())
    out, ckpt = _seed(tmp_path, [stampless, fresh])
    path = _write_xlsx_from_checkpoint(out, ckpt, {})
    wb = load_workbook(path)
    ws = wb["Matches"]
    headers = [c.value for c in ws[1]]
    job_col = headers.index("Job URL") + 1
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][job_col - 1] == "https://co.example.com/jobs/4"
    assert rows[-1][job_col - 1] == "https://co.example.com/jobs/3"
    wb.close()


def test_tracking_columns_have_dropdown_validation(tmp_path):
    out, ckpt = _seed(tmp_path, [_match(5, base_time().isoformat())])
    path = _write_xlsx_from_checkpoint(out, ckpt, {})
    wb = load_workbook(path)
    ws = wb["Matches"]
    headers = [c.value for c in ws[1]]
    dvs = ws.data_validations.dataValidation
    formulas = {str(dv.formula1) for dv in dvs}
    assert '"Yes,No"' in formulas
    assert any("not_fresher" in f for f in formulas)
    applied_col = headers.index("Applied") + 1
    feedback_col = headers.index("Feedback") + 1
    covered = {col for dv in dvs for rng in dv.sqref.ranges
               for col in (rng.min_col, rng.max_col)}
    assert applied_col in covered and feedback_col in covered
    wb.close()


def test_applied_jobs_leave_active_feed_into_applied_sheet(tmp_path):
    """Jobs marked Applied=Yes are excluded from the Matches feed and listed
    in their own sheet — tracked, but never cluttering the fresh list."""
    t = base_time()
    out, ckpt = _seed(tmp_path, [
        _match(1, (t - timedelta(days=1)).isoformat()),
        _match(2, t.isoformat()),
    ])
    path = _write_xlsx_from_checkpoint(
        out, ckpt, {}, applied_urls={"https://co.example.com/jobs/1"}
    )
    wb = load_workbook(path)
    headers = [c.value for c in wb["Matches"][1]]
    job_col = headers.index("Job URL") + 1
    feed_urls = {r[job_col - 1] for r in
                 wb["Matches"].iter_rows(min_row=2, values_only=True)}
    assert feed_urls == {"https://co.example.com/jobs/2"}

    assert "Applied" in wb.sheetnames
    ws_a = wb["Applied"]
    applied_urls = {r[job_col - 1] for r in
                    ws_a.iter_rows(min_row=2, values_only=True)}
    assert applied_urls == {"https://co.example.com/jobs/1"}
    applied_headers = [c.value for c in ws_a[1]]
    assert "First Seen" in applied_headers and "Notes" in applied_headers
    wb.close()


def test_no_applied_sheet_when_all_active(tmp_path):
    out, ckpt = _seed(tmp_path, [_match(6, base_time().isoformat())])
    path = _write_xlsx_from_checkpoint(out, ckpt, {})
    wb = load_workbook(path)
    assert "Applied" not in wb.sheetnames or \
        wb["Applied"].max_row == 1  # header-only at most
    wb.close()


def test_process_company_stamps_scraped_at(monkeypatch):
    """Every match carries a UTC scraped_at timestamp."""
    import saasnews_scraper as sc

    class FakeResp:
        status_code = 200
        text = '<html><a href="/careers">Careers</a></html>'
        url = "https://co.example.com/"

    class FakeSession:
        def get(self, url, timeout=None, allow_redirects=True, **kw):
            return FakeResp()

    monkeypatch.setattr(sc, "find_careers_page", lambda *a, **k: ("https://x/c", "found"))
    monkeypatch.setattr(
        sc, "fetch_jobs_for_careers_page",
        lambda *a, **k: ([sc.JobPosting(title="Python Backend Engineer",
                                        url="https://x/j1", location="Bengaluru",
                                        apply_url="https://x/j1")], "html"),
    )

    news = sc.NewsItem(news_url="https://example.com/news/1/",
                       headline="h", company_name="Co",
                       company_website="https://co.example.com")
    matches, rec, err = sc.process_company(news, FakeSession())
    assert err is None or err == {}
    assert matches
    for m in matches:
        parsed = datetime.fromisoformat(m.scraped_at)
        assert parsed.tzinfo is not None
