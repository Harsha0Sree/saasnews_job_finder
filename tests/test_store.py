"""Checkpoint persistence + xlsx rebuild.

Ported 1:1 from the original tests/run_tests.py custom runner (tmp_path
replaces tempfile.TemporaryDirectory).
"""
from openpyxl import load_workbook

from saasnews_scraper import (
    _append_checkpoint,
    _checkpoint_path,
    _load_checkpoint,
    _write_xlsx_from_checkpoint,
)


def _make_record(i: int) -> dict:
    return {
        "news_url": f"https://example.com/news/{i}/",
        "company_rec": {
            "company_name": f"Company{i}",
            "company_website": f"https://company{i}.com",
            "news_url": f"https://example.com/news/{i}/",
            "news_headline": f"Company{i} Raises Seed",
            "funding_round": "Seed",
            "funding_date": "July 2026",
            "lead_investor": "Test VC",
            "software_category": "AI",
            "careers_page_url": f"https://company{i}.com/careers",
            "careers_page_status": "found",
            "jobs_found": 5,
            "jobs_matched": 1,
            "error": "",
        },
        "matches": [
            {
                "company_name": f"Company{i}",
                "company_website": f"https://company{i}.com",
                "news_url": f"https://example.com/news/{i}/",
                "news_headline": f"Company{i} Raises Seed",
                "funding_round": "Seed",
                "funding_date": "July 2026",
                "software_category": "AI",
                "job_title": "Python Backend Engineer",
                "job_url": f"https://company{i}.com/jobs/1",
                "location": "Bengaluru",
                "apply_url": f"https://company{i}.com/jobs/1",
                "matched_keyword": "python",
                "match_category": "Backend",
                "confidence": "high",
                "location_basis": "india",
                "careers_page_url": f"https://company{i}.com/careers",
                "posted_date": "",
                "fresher_basis": "title_no_seniority",
                "fresher_level": "implicit",
            }
        ],
        "error": None,
    }


def test_empty_checkpoint_on_first_run(tmp_path):
    assert _load_checkpoint(str(tmp_path)) == {}


def test_checkpoint_roundtrip(tmp_path):
    out = str(tmp_path)
    for i in range(3):
        _append_checkpoint(out, _make_record(i))

    ckpt_path = _checkpoint_path(out)
    with open(ckpt_path) as f:
        lines = [l for l in f if l.strip()]
    assert len(lines) == 3

    ckpt = _load_checkpoint(out)
    assert len(ckpt) == 3
    expected_urls = {f"https://example.com/news/{i}/" for i in range(3)}
    assert set(ckpt.keys()) == expected_urls

    rec0 = ckpt["https://example.com/news/0/"]
    assert "company_rec" in rec0
    assert "matches" in rec0 and len(rec0["matches"]) == 1
    assert rec0["company_rec"]["company_name"] == "Company0"


def test_checkpoint_resume_appends(tmp_path):
    out = str(tmp_path)
    for i in range(3):
        _append_checkpoint(out, _make_record(i))
    rec3 = _make_record(3)
    rec3["company_rec"].update({
        "careers_page_status": "not_found", "jobs_found": 0,
        "jobs_matched": 0, "careers_page_url": "",
    })
    rec3["matches"] = []
    _append_checkpoint(out, rec3)
    assert len(_load_checkpoint(out)) == 4


def test_corrupt_lines_skipped(tmp_path):
    out = str(tmp_path)
    _append_checkpoint(out, _make_record(0))
    with open(_checkpoint_path(out), "a") as f:
        f.write("{not json at all\n")
    _append_checkpoint(out, _make_record(1))
    assert len(_load_checkpoint(out)) == 2


def test_run_summary_reports_real_discovered_count(tmp_path):
    """The 'Total News Articles Discovered' figure must be the real count
    passed by the caller — never pages×12 arithmetic, and zero stays zero."""
    out = str(tmp_path)
    _append_checkpoint(out, _make_record(0))
    ckpt = _load_checkpoint(out)
    out_path = _write_xlsx_from_checkpoint(
        out, ckpt, {"news_pages": 0}, news_total=137
    )
    wb = load_workbook(out_path)
    ws = wb["Run Summary"]
    assert ws.cell(row=2, column=1).value == "Total News Articles Discovered"
    assert ws.cell(row=2, column=2).value == 137
    wb.close()


def test_run_summary_zero_discovered_stays_zero(tmp_path):
    out = str(tmp_path)
    _append_checkpoint(out, _make_record(0))
    ckpt = _load_checkpoint(out)
    out_path = _write_xlsx_from_checkpoint(
        out, ckpt, {"news_pages": 0}, news_total=0
    )
    wb = load_workbook(out_path)
    # 0 discovered (not the 1 checkpointed company masquerading as articles).
    assert wb["Run Summary"].cell(row=2, column=2).value == 0
    wb.close()


def test_xlsx_rebuilt_from_checkpoint(tmp_path):
    out = str(tmp_path)
    for i in range(4):
        rec = _make_record(i)
        if i == 3:
            rec["matches"] = []
            rec["company_rec"].update({
                "careers_page_status": "not_found", "jobs_found": 0,
                "jobs_matched": 0, "careers_page_url": "",
            })
        _append_checkpoint(out, rec)

    ckpt = _load_checkpoint(out)
    out_path = _write_xlsx_from_checkpoint(
        out, ckpt, {"news_pages": 50, "concurrency": 10}
    )
    wb = load_workbook(out_path)
    assert "Matches" in wb.sheetnames
    assert "All Companies" in wb.sheetnames
    assert "Run Summary" in wb.sheetnames
    ws = wb["Matches"]
    # 3 matches (Company0,1,2 each have 1 match; Company3 has 0) + header
    assert ws.max_row == 4
    ws = wb["All Companies"]
    # 4 companies + header
    assert ws.max_row == 5
    wb.close()
